from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from pathlib import Path

from docx import Document

from experiments.common.metadata import ROOT
from experiments.cv.update_external_cvs import CN_AFTER as CN_V1, CN_FILE
from experiments.cv.update_external_cvs import EN_AFTER as EN_V1, EN_FILE


EN_V2 = (
    "Benchmarked the structured selective-LLM architecture against a compact full-context "
    "LLM baseline on 160 synthetic ordering conversations (868 turns) using an offline "
    "deterministic decision replay, reducing model-specific estimated input-token usage by "
    "92.3% and simulated LLM calls by 94.0% while preserving expected-path performance and "
    "zero blocking safety violations in the evaluation set."
)
CN_V2 = (
    "在 160 组合成订餐会话（868 轮）上，以离线确定性决策回放对结构化 selective-LLM "
    "架构与紧凑完整上下文 LLM 基线进行成对评测；在保持预期路径表现及阻断级安全违规为 0 "
    "的情况下，模型特定 tokenizer 估算的输入 Token 使用量降低 92.3%，模拟 LLM 调用次数降低 94.0%。"
)

ELIGIBLE_STATUSES = {"VERIFIED_FOR_CV", "CV_ELIGIBLE_WITH_QUALIFIER"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_exact_paragraph(path: Path, before: str, after: str) -> dict[str, object]:
    before_sha = sha256(path)
    document = Document(path)
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == before]
    if len(matches) != 1:
        raise ValueError(f"expected exactly one v1 metric paragraph in {path.name}, found {len(matches)}")
    paragraph = matches[0]
    if len(paragraph.runs) != 1:
        raise ValueError(f"target paragraph in {path.name} has {len(paragraph.runs)} runs")
    style_name = paragraph.style.name
    paragraph.runs[0].text = after
    document.save(path)

    verified = Document(path)
    paragraphs = [item.text for item in verified.paragraphs]
    if paragraphs.count(after) != 1 or before in paragraphs:
        raise AssertionError(f"metric replacement was not exact in {path.name}")
    if next(item for item in verified.paragraphs if item.text == after).style.name != style_name:
        raise AssertionError(f"paragraph style changed in {path.name}")
    return {
        "path": str(path.resolve()),
        "before_sha256": before_sha,
        "after_sha256": sha256(path),
        "paragraph_style": style_name,
        "before": before,
        "after": after,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cv-dir", type=Path, required=True)
    parser.add_argument("--backup-dir", type=Path, required=True)
    parser.add_argument(
        "--analysis",
        type=Path,
        default=ROOT / "experiments" / "results" / "token_efficiency_v2" / "analysis.json",
    )
    parser.add_argument(
        "--audit-output",
        type=Path,
        default=ROOT / "experiments" / "results" / "token_efficiency_v2" / "cv_update_v2.json",
    )
    args = parser.parse_args()

    analysis = json.loads(args.analysis.read_text(encoding="utf-8"))
    status = analysis["cv_review"]["status"]
    if status not in ELIGIBLE_STATUSES:
        raise RuntimeError(f"CV update blocked by metric status: {status}")
    if analysis["cv_review"]["allowed_english_wording"] != EN_V2:
        raise AssertionError("English CV wording differs from audited wording")
    if analysis["cv_review"]["allowed_chinese_wording"] != CN_V2:
        raise AssertionError("Chinese CV wording differs from audited wording")

    source_files = [
        args.cv_dir / EN_FILE,
        args.cv_dir / EN_FILE.replace(".docx", ".pdf"),
        args.cv_dir / CN_FILE,
        args.cv_dir / CN_FILE.replace(".docx", ".pdf"),
    ]
    if args.backup_dir.exists():
        raise FileExistsError(f"refusing to overwrite backup directory: {args.backup_dir}")
    args.backup_dir.mkdir(parents=True)
    backup_files = []
    for source in source_files:
        destination = args.backup_dir / source.name
        shutil.copy2(source, destination)
        backup_files.append(
            {"source": str(source.resolve()), "backup": str(destination.resolve()), "sha256": sha256(destination)}
        )

    updates = [
        replace_exact_paragraph(args.cv_dir / EN_FILE, EN_V1, EN_V2),
        replace_exact_paragraph(args.cv_dir / CN_FILE, CN_V1, CN_V2),
    ]
    payload = {
        "source_metric": "experiments/reports/resume_metrics_review_v2.md",
        "metric_status": status,
        "backup_files": backup_files,
        "numeric_consistency": {
            "conversations": 160,
            "turns": 868,
            "input_token_reduction_percent": 92.3,
            "simulated_call_reduction_percent": 94.0,
        },
        "qualifiers": [
            "synthetic conversations",
            "offline deterministic decision replay",
            "model-specific estimated input tokens",
            "simulated LLM calls",
        ],
        "excluded_metrics": [
            "actual API billing",
            "live-model semantic accuracy",
            "local/replay latency",
            "barge-in/recovery",
            "real speech/phone/user/restaurant outcomes",
        ],
        "updates": updates,
    }
    args.audit_output.parent.mkdir(parents=True, exist_ok=True)
    args.audit_output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

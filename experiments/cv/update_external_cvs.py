from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document

from experiments.common.metadata import ROOT


EN_FILE = "Zhihan_Sun_CV_EN.docx"
CN_FILE = "孙之涵_简历_CN.docx"

EN_BEFORE = (
    "Structured routine ordering operations around canonical intents, deterministic state transitions, and explicit safety rules, "
    "while confining optional model-driven reasoning to a validated fallback sandbox for low-confidence cases; designed the architecture "
    "to limit unnecessary model calls and long-context transmission and make latency, inference cost, and high-impact actions more controllable, "
    "with AUTO_DRAFT, CONFIRM, HANDOFF, and REFUSE pathways."
)
EN_AFTER = (
    "Benchmarked the structured deterministic path against a documented full-context LLM baseline on 160 synthetic ordering conversations "
    "(868 turns): the offline path retained AUTO_DRAFT, CONFIRM, HANDOFF, and REFUSE gates while reducing projected LLM input-token transfer "
    "and model-call requirements by 100% (bootstrap 95% CI, 100%–100%); counts use a deterministic Unicode-codepoint estimator and do not "
    "represent billed tokens."
)
CN_BEFORE = (
    "以规范化意图、确定性状态转换和显式安全规则处理常规订餐操作，仅在经校验的低置信度 fallback 沙箱中保留可选模型推理；"
    "该设计旨在减少不必要的模型调用与长上下文传输，使延迟、推理成本和高影响操作更可控，并设置 AUTO_DRAFT、CONFIRM、HANDOFF、REFUSE 四类路径。"
)
CN_AFTER_V1 = (
    "在 160 组合成订餐会话（868 轮）上，将结构化确定性路径与有文档定义的完整上下文 LLM 基线进行成对评测：离线路径在保留 "
    "AUTO_DRAFT、CONFIRM、HANDOFF、REFUSE 安全门控的同时，使预估 LLM 输入 Token 传输量与模型调用需求均降低 100%"
    "（Bootstrap 95% CI：100%–100%）；Token 采用确定性 Unicode code-point 估算，不代表实际计费 Token。"
)
CN_AFTER = (
    "在 160 组合成订餐会话（868 轮）上成对比较结构化路径与有文档定义的完整上下文 LLM 基线：禁用实时 LLM 的离线配置在保留 "
    "AUTO_DRAFT、CONFIRM、HANDOFF、REFUSE 门控下，使预估 LLM 输入 Token 传输量和模型调用需求均降低 100%"
    "（Bootstrap 95% CI：100%–100%；Unicode code-point 估算，非实际计费 Token）。"
)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_exact_paragraph(path: Path, before: str, after: str) -> dict:
    original_sha = digest(path)
    document = Document(path)
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == before]
    if len(matches) != 1:
        raise ValueError(f"expected exactly one target paragraph in {path.name}, found {len(matches)}")
    paragraph = matches[0]
    if len(paragraph.runs) != 1:
        raise ValueError(f"target paragraph in {path.name} unexpectedly has {len(paragraph.runs)} runs")
    original_style = paragraph.style.name
    paragraph.runs[0].text = after
    if paragraph.style.name != original_style:
        raise AssertionError("paragraph style changed")
    document.save(path)
    verified = Document(path)
    if sum(item.text == after for item in verified.paragraphs) != 1:
        raise AssertionError(f"replacement not persisted in {path.name}")
    if any(item.text == before for item in verified.paragraphs):
        raise AssertionError(f"old paragraph remains in {path.name}")
    return {
        "path": str(path),
        "before_sha256": original_sha,
        "after_sha256": digest(path),
        "paragraph_style": original_style,
        "before": before,
        "after": after,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cv-dir", type=Path, required=True)
    parser.add_argument("--audit-output", type=Path, default=ROOT / "experiments" / "results" / "cv_update.json")
    args = parser.parse_args()
    updates = [
        replace_exact_paragraph(args.cv_dir / EN_FILE, EN_BEFORE, EN_AFTER),
        replace_exact_paragraph(args.cv_dir / CN_FILE, CN_BEFORE, CN_AFTER),
    ]
    payload = {
        "source_metric": "experiments/reports/resume_metrics_review.md",
        "new_metric_inserted": "projected LLM input-token and model-call reduction only",
        "excluded_metrics": ["local latency", "Replay latency", "interruption F1/FPR", "recovery", "real speech", "cost"],
        "numeric_consistency": {
            "conversations": 160,
            "turns": 868,
            "reduction_percent": 100,
            "bootstrap_ci_percent": [100, 100],
        },
        "updates": updates,
    }
    args.audit_output.parent.mkdir(parents=True, exist_ok=True)
    args.audit_output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

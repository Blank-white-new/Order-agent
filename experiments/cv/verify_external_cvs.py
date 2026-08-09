from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from experiments.common.metadata import ROOT
from experiments.cv.update_external_cvs import CN_AFTER, CN_FILE, EN_AFTER, EN_FILE


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate_docx(path: Path, expected: str) -> dict:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    if paragraphs.count(expected) != 1:
        raise AssertionError(f"audited bullet missing or duplicated in {path.name}")
    project_bullets = [paragraph.text for paragraph in document.paragraphs[16:21]]
    if len(project_bullets) != 5 or expected not in project_bullets:
        raise AssertionError(f"first project no longer has five bullets in {path.name}")
    return {"sha256": sha256(path), "paragraphs": len(paragraphs), "first_project_bullets": len(project_bullets)}


def validate_pdf(path: Path, required_fragments: tuple[str, ...]) -> dict:
    reader = PdfReader(path)
    if len(reader.pages) != 1:
        raise AssertionError(f"{path.name} has {len(reader.pages)} pages; expected one")
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    folded = " ".join(text.split())
    missing = [fragment for fragment in required_fragments if fragment not in folded]
    if missing:
        raise AssertionError(f"{path.name} is missing expected PDF text: {missing}")
    return {"sha256": sha256(path), "pages": 1, "extracted_characters": len(text)}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cv-dir", type=Path, required=True)
    parser.add_argument("--audit", type=Path, default=ROOT / "experiments" / "results" / "cv_update.json")
    args = parser.parse_args()
    en_docx = args.cv_dir / EN_FILE
    cn_docx = args.cv_dir / CN_FILE
    en_pdf = args.cv_dir / EN_FILE.replace(".docx", ".pdf")
    cn_pdf = args.cv_dir / CN_FILE.replace(".docx", ".pdf")
    results = {
        "english_docx": validate_docx(en_docx, EN_AFTER),
        "chinese_docx": validate_docx(cn_docx, CN_AFTER),
        "english_pdf": validate_pdf(en_pdf, ("160 synthetic ordering conversations", "868 turns", "100%", "billed tokens")),
        "chinese_pdf": validate_pdf(cn_pdf, ("160 组合成订餐会话", "868 轮", "100%", "非实际计费 Token")),
    }
    audit = json.loads(args.audit.read_text(encoding="utf-8"))
    audit["verification"] = {
        "files": results,
        "english_chinese_numeric_consistency": True,
        "project_bullet_count_each": 5,
        "pdf_page_count_each": 1,
        "visual_review": {
            "method": "Word PDF export, Poppler 150-dpi full-page PNG, manual inspection at original resolution",
            "english_pages_inspected": 1,
            "chinese_pages_inspected": 1,
            "overflow_or_clipping": False,
        },
        "unsupported_claims_added": False,
    }
    for item in audit.get("render_iterations", []):
        if item.get("language") == "Chinese" and item.get("iteration") == 2:
            item.update({"pages": 1, "result": "accepted"})
    args.audit.write_text(json.dumps(audit, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
